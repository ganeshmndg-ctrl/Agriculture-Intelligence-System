from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# DAILY INTELLIGENCE WORD REPORT GENERATOR v1.0
# ============================================================
#
# Reads:
#   data/daily_intelligence_v1/
#       DAILY_INTELLIGENCE_DATABASE_v1.json
#
# Creates:
#   data/daily_reports/
#       YYYY-MM-DD_Agriculture_Intelligence.docx
#
# Important:
#   - No Gemini/API calls
#   - No new facts
#   - No question generation
#   - Uses only the verified Daily Intelligence database
# ============================================================


BASE_DIR = Path(__file__).resolve().parent

INPUT_FILE = (
    BASE_DIR
    / "data"
    / "daily_intelligence_v1"
    / "DAILY_INTELLIGENCE_DATABASE_v1.json"
)

OUTPUT_DIR = (
    BASE_DIR
    / "data"
    / "daily_reports"
)


# ============================================================
# CATEGORY ORDER
# ============================================================

CATEGORY_ORDER = [
    "CORE_AGRICULTURE",
    "BANKING_FINANCE",
    "INTERNATIONAL_TRADE",
    "MACRO_ECONOMY",
    "GENERAL_BUSINESS",
]

CATEGORY_LABELS = {
    "CORE_AGRICULTURE": "Core Agriculture",
    "BANKING_FINANCE": "Banking & Finance",
    "INTERNATIONAL_TRADE": "International Trade",
    "MACRO_ECONOMY": "Macro Economy",
    "GENERAL_BUSINESS": "General Business",
}


# ============================================================
# BASIC HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""

    return str(value).strip()


def as_list(value):
    if isinstance(value, list):
        return value

    if value is None:
        return []

    return [value]


def load_database():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Daily Intelligence database not found:\n{INPUT_FILE}"
        )

    import json

    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


# ============================================================
# WORD FORMATTING HELPERS
# ============================================================

def set_cell_shading(cell, fill):
    """Set table-cell background colour."""

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(
        qn("w:shd")
    )

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(
        qn("w:fill"),
        fill
    )


def set_cell_text(cell, text, bold=False):
    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        clean(text)
    )

    run.bold = bold
    run.font.size = Pt(9)

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


def add_horizontal_line(paragraph):
    """
    Add a simple bottom border to a paragraph.
    """

    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(
        qn("w:val"),
        "single"
    )

    bottom.set(
        qn("w:sz"),
        "6"
    )

    bottom.set(
        qn("w:space"),
        "1"
    )

    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_heading(doc, text, level=1):

    paragraph = doc.add_heading(
        text,
        level=level
    )

    return paragraph


def add_bullet(doc, text, level=0):

    text = clean(text)

    if not text:
        return

    style = (
        "List Bullet"
        if level == 0
        else "List Bullet 2"
    )

    paragraph = doc.add_paragraph(
        style=style
    )

    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text)
    run.font.size = Pt(10)


def add_number(doc, text):

    text = clean(text)

    if not text:
        return

    paragraph = doc.add_paragraph(
        style="List Number"
    )

    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text)
    run.font.size = Pt(10)


def add_label_paragraph(doc, label, text):

    text = clean(text)

    if not text:
        return

    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(6)

    run = paragraph.add_run(
        label
    )

    run.bold = True
    run.font.size = Pt(10)

    run2 = paragraph.add_run(
        text
    )

    run2.font.size = Pt(10)


def add_list_section(doc, heading, values):

    values = as_list(values)

    if not values:
        return

    add_heading(
        doc,
        heading,
        level=3
    )

    for value in values:

        if isinstance(value, dict):

            # Try to present dictionary information
            # without losing any useful content.

            for key, val in value.items():

                if isinstance(val, list):
                    val = ", ".join(
                        clean(x)
                        for x in val
                    )

                if clean(val):
                    add_bullet(
                        doc,
                        f"{key}: {val}"
                    )

        else:

            if clean(value):
                add_bullet(
                    doc,
                    clean(value)
                )


# ============================================================
# DOCUMENT STYLE
# ============================================================

def configure_document(doc):

    section = doc.sections[0]

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles

    normal = styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(10)

    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [
        ("Title", 24),
        ("Heading 1", 17),
        ("Heading 2", 14),
        ("Heading 3", 11),
    ]:

        style = styles[style_name]

        style.font.name = "Aptos"
        style.font.size = Pt(size)


# ============================================================
# COVER / TITLE
# ============================================================

def add_document_header(
    doc,
    database
):

    date_values = database.get(
        "dates",
        {}
    )

    # Usually the database contains one newspaper
    # date for a daily report.
    if date_values:
        report_date = next(
            iter(date_values.keys())
        )
    else:
        report_date = datetime.now().strftime(
            "%Y-%m-%d"
        )

    try:

        parsed_date = datetime.strptime(
            report_date,
            "%Y-%m-%d"
        )

        display_date = parsed_date.strftime(
            "%d %B %Y"
        )

    except ValueError:

        display_date = report_date

    title = doc.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        "AGRICULTURE & BUSINESS\n"
        "INTELLIGENCE"
    )

    run.bold = True
    run.font.size = Pt(25)

    subtitle = doc.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        f"Daily Intelligence Report\n"
        f"{display_date}"
    )

    run.font.size = Pt(13)

    newspapers = database.get(
        "newspapers",
        {}
    )

    newspaper_names = list(
        newspapers.keys()
    )

    if newspaper_names:

        source_paragraph = doc.add_paragraph()

        source_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = source_paragraph.add_run(
            "Sources: "
            + " | ".join(
                newspaper_names
            )
        )

        run.italic = True
        run.font.size = Pt(10)

    doc.add_paragraph()

    intro = doc.add_paragraph()

    intro.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = intro.add_run(
        "Prepared from the verified "
        "Daily Intelligence Database"
    )

    run.font.size = Pt(9)
    run.italic = True

    doc.add_page_break()


# ============================================================
# INTELLIGENCE SNAPSHOT
# ============================================================

def add_snapshot(
    doc,
    database
):

    add_heading(
        doc,
        "1. Daily Intelligence Snapshot",
        level=1
    )

    summary = database.get(
        "summary",
        {}
    )

    table = doc.add_table(
        rows=0,
        cols=2
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = "Table Grid"

    rows = [
        (
            "Total analysed stories",
            summary.get(
                "total_records",
                ""
            )
        ),
        (
            "Government Intelligence records",
            summary.get(
                "government_intelligence_records",
                ""
            )
        ),
        (
            "Records with connections",
            summary.get(
                "records_with_connections",
                ""
            )
        ),
        (
            "Records with key numbers",
            summary.get(
                "records_with_key_numbers",
                ""
            )
        ),
        (
            "Records with interpretations",
            summary.get(
                "records_with_interpretations",
                ""
            )
        ),
    ]

    for label, value in rows:

        cells = table.add_row().cells

        set_cell_text(
            cells[0],
            label,
            bold=True
        )

        set_cell_text(
            cells[1],
            value
        )

        set_cell_shading(
            cells[0],
            "EAF2F8"
        )

    doc.add_paragraph()

    # Categories
    categories = summary.get(
        "categories",
        {}
    )

    if categories:

        add_heading(
            doc,
            "Coverage by category",
            level=2
        )

        for category in CATEGORY_ORDER:

            count = categories.get(
                category
            )

            if count is not None:

                add_bullet(
                    doc,
                    f"{CATEGORY_LABELS.get(category, category)}: "
                    f"{count}"
                )

        # Include unexpected categories too
        for category, count in categories.items():

            if category not in CATEGORY_ORDER:

                add_bullet(
                    doc,
                    f"{category}: {count}"
                )

    doc.add_paragraph()


# ============================================================
# SOURCE INFORMATION
# ============================================================

def add_source_information(
    doc,
    record
):

    source = record.get(
        "source",
        {}
    )

    if not isinstance(source, dict):
        source = {}

    add_heading(
        doc,
        "Source",
        level=3
    )

    table = doc.add_table(
        rows=0,
        cols=2
    )

    table.style = "Table Grid"

    source_rows = [
        (
            "Newspaper",
            source.get(
                "newspaper",
                record.get(
                    "newspaper",
                    ""
                )
            )
        ),
        (
            "Date",
            source.get(
                "date",
                record.get(
                    "date",
                    ""
                )
            )
        ),
        (
            "Page",
            source.get(
                "page",
                record.get(
                    "page",
                    ""
                )
            )
        ),
        (
            "Original headline",
            source.get(
                "headline",
                record.get(
                    "headline",
                    ""
                )
            )
        ),
        (
            "Source PDF",
            source.get(
                "source_pdf",
                ""
            )
        ),
    ]

    for label, value in source_rows:

        if not clean(value):
            continue

        cells = table.add_row().cells

        set_cell_text(
            cells[0],
            label,
            bold=True
        )

        set_cell_text(
            cells[1],
            value
        )

        set_cell_shading(
            cells[0],
            "F2F2F2"
        )

    doc.add_paragraph()


# ============================================================
# GOVERNMENT INTELLIGENCE
# ============================================================

def add_government_intelligence(
    doc,
    record
):

    government = as_list(
        record.get(
            "government_intelligence"
        )
    )

    if not government:
        return

    add_heading(
        doc,
        "Government Intelligence",
        level=3
    )

    for initiative in government:

        if not isinstance(
            initiative,
            dict
        ):
            add_bullet(
                doc,
                initiative
            )
            continue

        name = clean(
            initiative.get(
                "initiative_name"
            )
        )

        if name:

            paragraph = doc.add_paragraph()

            run = paragraph.add_run(
                name
            )

            run.bold = True
            run.font.size = Pt(10)

        fields = [
            (
                "Type",
                initiative.get(
                    "initiative_type"
                )
            ),
            (
                "Event",
                initiative.get(
                    "event_type"
                )
            ),
            (
                "Details",
                initiative.get(
                    "details"
                )
            ),
        ]

        for label, value in fields:

            if clean(value):

                add_label_paragraph(
                    doc,
                    f"{label}: ",
                    value
                )

        source_details = []

        if clean(
            initiative.get("newspaper")
        ):
            source_details.append(
                initiative.get(
                    "newspaper"
                )
            )

        if clean(
            initiative.get("date")
        ):
            source_details.append(
                initiative.get(
                    "date"
                )
            )

        if clean(
            initiative.get("page")
        ):
            source_details.append(
                f"Page {initiative.get('page')}"
            )

        if source_details:

            add_label_paragraph(
                doc,
                "Source: ",
                " | ".join(
                    source_details
                )
            )

        doc.add_paragraph()


# ============================================================
# CONNECTIONS
# ============================================================

def add_connections(
    doc,
    record
):

    connections = as_list(
        record.get(
            "connections"
        )
    )

    if not connections:
        return

    add_heading(
        doc,
        "Connections to Other News",
        level=3
    )

    for connection in connections:

        if not isinstance(
            connection,
            dict
        ):
            add_bullet(
                doc,
                connection
            )
            continue

        entities = connection.get(
            "entities",
            []
        )

        if entities:

            paragraph = doc.add_paragraph()

            run = paragraph.add_run(
                "Entities: "
            )

            run.bold = True

            paragraph.add_run(
                " → ".join(
                    clean(x)
                    for x in entities
                )
            )

        relationship = clean(
            connection.get(
                "relationship"
            )
        )

        if relationship:

            add_label_paragraph(
                doc,
                "Relationship: ",
                relationship
            )

        chain = connection.get(
            "chain",
            []
        )

        if chain:

            add_label_paragraph(
                doc,
                "Chain: ",
                " → ".join(
                    clean(x)
                    for x in chain
                )
            )

        explanation = clean(
            connection.get(
                "explanation"
            )
        )

        if explanation:

            add_label_paragraph(
                doc,
                "Explanation: ",
                explanation
            )

        confidence = clean(
            connection.get(
                "confidence"
            )
        )

        if confidence:

            add_label_paragraph(
                doc,
                "Confidence: ",
                confidence
            )

        doc.add_paragraph()


# ============================================================
# INDIVIDUAL ARTICLE
# ============================================================

def add_article(
    doc,
    record,
    article_number
):

    category = clean(
        record.get(
            "primary_category"
        )
    )

    category_label = CATEGORY_LABELS.get(
        category,
        category
    )

    headline = clean(
        record.get(
            "headline"
        )
    )

    add_heading(
        doc,
        f"{article_number}. {headline}",
        level=2
    )

    metadata = []

    if category_label:
        metadata.append(
            category_label
        )

    if clean(
        record.get(
            "intelligence_type"
        )
    ):
        metadata.append(
            clean(
                record.get(
                    "intelligence_type"
                )
            )
        )

    if clean(
        record.get(
            "agriculture_relevance"
        )
    ):
        metadata.append(
            "Agriculture relevance: "
            + clean(
                record.get(
                    "agriculture_relevance"
                )
            )
        )

    if metadata:

        paragraph = doc.add_paragraph()

        run = paragraph.add_run(
            " | ".join(metadata)
        )

        run.italic = True
        run.font.size = Pt(9)

    # --------------------------------------------------------
    # What happened
    # --------------------------------------------------------

    add_label_paragraph(
        doc,
        "What happened: ",
        record.get(
            "what_happened"
        )
    )

    # --------------------------------------------------------
    # Key facts
    # --------------------------------------------------------

    add_list_section(
        doc,
        "Key facts",
        record.get(
            "key_facts"
        )
    )

    # --------------------------------------------------------
    # Reported facts
    # --------------------------------------------------------

    add_list_section(
        doc,
        "Reported facts",
        record.get(
            "reported_facts"
        )
    )

    # --------------------------------------------------------
    # Key numbers
    # --------------------------------------------------------

    add_list_section(
        doc,
        "Key numbers",
        record.get(
            "key_numbers"
        )
    )

    # --------------------------------------------------------
    # Why it matters
    # --------------------------------------------------------

    add_label_paragraph(
        doc,
        "Why it matters: ",
        record.get(
            "why_it_matters"
        )
    )

    # --------------------------------------------------------
    # Agriculture impact
    # --------------------------------------------------------

    add_label_paragraph(
        doc,
        "Agriculture impact: ",
        record.get(
            "agriculture_impact"
        )
    )

    # --------------------------------------------------------
    # Agriculture relevance reason
    # --------------------------------------------------------

    add_label_paragraph(
        doc,
        "Agriculture relevance: ",
        record.get(
            "agriculture_relevance_reason"
        )
    )

    # --------------------------------------------------------
    # Interpretations
    # --------------------------------------------------------

    add_list_section(
        doc,
        "Interpretations",
        record.get(
            "interpretations"
        )
    )

    # --------------------------------------------------------
    # Topics
    # --------------------------------------------------------

    add_list_section(
        doc,
        "Topics",
        record.get(
            "topics"
        )
    )

    # --------------------------------------------------------
    # Important terms
    # --------------------------------------------------------

    important_terms = as_list(
        record.get(
            "important_terms"
        )
    )

    if important_terms:

        paragraph = doc.add_paragraph()

        run = paragraph.add_run(
            "Important terms: "
        )

        run.bold = True

        paragraph.add_run(
            " | ".join(
                clean(x)
                for x in important_terms
                if clean(x)
            )
        )

    # --------------------------------------------------------
    # Government intelligence
    # --------------------------------------------------------

    add_government_intelligence(
        doc,
        record
    )

    # --------------------------------------------------------
    # Connections
    # --------------------------------------------------------

    add_connections(
        doc,
        record
    )

    # --------------------------------------------------------
    # Article text
    # --------------------------------------------------------

    article_text = clean(
        record.get(
            "article_text"
        )
    )

    if article_text:

        add_heading(
            doc,
            "Original Article Text",
            level=3
        )

        paragraph = doc.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(8)

        run = paragraph.add_run(
            article_text
        )

        run.font.size = Pt(9)

    # --------------------------------------------------------
    # Source
    # --------------------------------------------------------

    add_source_information(
        doc,
        record
    )

    # Separator
    paragraph = doc.add_paragraph()

    add_horizontal_line(
        paragraph
    )


# ============================================================
# BUILD WORD DOCUMENT
# ============================================================

def build_word_document(database):

    records = database.get(
        "records",
        []
    )

    if not records:
        raise ValueError(
            "Daily Intelligence database contains no records."
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    dates = database.get(
        "summary",
        {}
    ).get(
        "dates",
        {}
    )

    if dates:

        report_date = next(
            iter(dates.keys())
        )

    else:

        report_date = datetime.now().strftime(
            "%Y-%m-%d"
        )

    output_file = (
        OUTPUT_DIR
        / f"{report_date}_Agriculture_Intelligence.docx"
    )

    doc = Document()

    configure_document(
        doc
    )

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    add_document_header(
        doc,
        database
    )

    # --------------------------------------------------------
    # Snapshot
    # --------------------------------------------------------

    add_snapshot(
        doc,
        database
    )

    # --------------------------------------------------------
    # Articles by category
    # --------------------------------------------------------

    article_number = 1

    for category in CATEGORY_ORDER:

        category_records = [
            record
            for record in records
            if clean(
                record.get(
                    "primary_category"
                )
            ) == category
        ]

        if not category_records:
            continue

        add_heading(
            doc,
            CATEGORY_LABELS.get(
                category,
                category
            ),
            level=1
        )

        for record in category_records:

            add_article(
                doc,
                record,
                article_number
            )

            article_number += 1

    # --------------------------------------------------------
    # Unexpected categories
    # --------------------------------------------------------

    known_categories = set(
        CATEGORY_ORDER
    )

    unexpected = [
        record
        for record in records
        if clean(
            record.get(
                "primary_category"
            )
        ) not in known_categories
    ]

    if unexpected:

        add_heading(
            doc,
            "Other Intelligence",
            level=1
        )

        for record in unexpected:

            add_article(
                doc,
                record,
                article_number
            )

            article_number += 1

    # --------------------------------------------------------
    # Final note
    # --------------------------------------------------------

    doc.add_page_break()

    add_heading(
        doc,
        "End of Daily Intelligence Report",
        level=1
    )

    paragraph = doc.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        "This document was generated from the "
        "verified Daily Intelligence Database. "
        "No additional facts were introduced "
        "during document generation."
    )

    run.font.size = Pt(9)
    run.italic = True

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    doc.save(
        output_file
    )

    return output_file


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print(
        "DAILY INTELLIGENCE WORD REPORT GENERATOR v1.0"
    )
    print("=" * 70)
    print()

    print(
        f"Input:\n{INPUT_FILE}"
    )

    print()

    database = load_database()

    print(
        f"Records loaded: "
        f"{len(database.get('records', []))}"
    )

    print()

    output_file = build_word_document(
        database
    )

    print("=" * 70)
    print("WORD REPORT CREATED")
    print("=" * 70)

    print(
        f"Output:\n{output_file}"
    )

    print()

    print(
        "The report contains the analysed "
        "newspaper intelligence, government "
        "intelligence, connections, key facts, "
        "agriculture impact and source information."
    )

    print("=" * 70)


if __name__ == "__main__":
    main()